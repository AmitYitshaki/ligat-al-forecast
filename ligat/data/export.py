"""Export validated clean datasets for inspection in Excel."""

import hashlib
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import pandas as pd  # type: ignore[import-untyped]
import yaml
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from ligat.data.repository import Repository

DATASETS = (
    "matches",
    "fixtures",
    "round_tables",
    "final_tables",
    "game_details",
    "goals",
    "red_cards",
    "coaches",
)


def _fingerprints(repo: Repository, teams_path: Path) -> dict[str, str]:
    return {
        name: hashlib.sha256((repo.root / f"{name}.parquet").read_bytes()).hexdigest()
        for name in DATASETS
    } | {"teams_yaml": hashlib.sha256(teams_path.read_bytes()).hexdigest()}


def _versioned_path(repo: Repository, hashes: dict[str, str], major: bool) -> tuple[Path, bool]:
    versions = []
    for path in repo.root.glob("ligat_al_data_*.xlsx"):
        match = re.fullmatch(r"ligat_al_data_(\d+)\.(\d+)\.xlsx", path.name)
        if match:
            versions.append((tuple(map(int, match.groups())), path))
    if versions:
        (current_major, minor), latest = max(versions)
        metadata = latest.with_suffix(".json")
        if not major and metadata.exists():
            saved = json.loads(metadata.read_text(encoding="utf-8"))
            if saved.get("source_sha256") == hashes:
                if hashlib.sha256(latest.read_bytes()).hexdigest() != saved["excel_sha256"]:
                    raise ValueError(
                        "Versioned Excel file changed; it cannot be silently overwritten"
                    )
                return latest, True
        version = (current_major + 1, 0) if major else (current_major, minor + 1)
    else:
        version = (1, 0) if major else (0, 3)
    return repo.root / f"ligat_al_data_{version[0]}.{version[1]}.xlsx", False


def _excel_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        # Excel cannot store timezones. Keep the Israeli wall clock, never shift to UTC.
        return value.tz_convert("Asia/Jerusalem").tz_localize(None).round("ms").to_pydatetime()
    return value.item() if hasattr(value, "item") else value


def export_excel(
    repository: Repository | None = None,
    output: Path | None = None,
    *,
    major: bool = False,
    teams_path: Path = Path("config/teams.yaml"),
) -> Path:
    """Mirror all eight datasets, creating an immutable numbered workbook by default."""
    if output is not None:
        if "raw" in (part.casefold() for part in output.resolve().parts):
            raise ValueError("Excel export cannot write into raw storage")
        if output.suffix.lower() != ".xlsx":
            raise ValueError("Excel output must have an .xlsx suffix")
        if re.fullmatch(r"ligat_al_data_\d+\.\d+\.xlsx", output.name) and output.exists():
            raise FileExistsError("An existing numbered Excel version cannot be overwritten")
    repo = repository if repository is not None else Repository()
    frames = {name: repo.read(name) for name in DATASETS}
    teams = yaml.safe_load(teams_path.read_text(encoding="utf-8"))
    names = {row["team_id"]: row["name_he"] for row in teams["teams"]}
    for name, frame in frames.items():
        for column in ("home_team_id", "away_team_id", "team_id"):
            if column not in frame:
                continue
            missing = set(frame[column]) - set(names)
            if missing:
                raise ValueError(f"Missing canonical team names in {name}: {sorted(missing)}")
            label = column.removesuffix("_id") + "_name_he"
            frame.insert(frame.columns.get_loc(column) + 1, label, frame[column].map(names))
    hashes = _fingerprints(repo, teams_path) if output is None else None
    if output is None:
        assert hashes is not None
        output, cached = _versioned_path(repo, hashes, major)
        if cached:
            return output
    if "raw" in (part.casefold() for part in output.resolve().parts):
        raise ValueError("Excel export cannot write into raw storage")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Excel output must have an .xlsx suffix")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in frames.items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(frame.columns))
        for record in frame.itertuples(index=False, name=None):
            sheet.append([_excel_value(value) for value in record])
        sheet.freeze_panes = "A2"
        sheet.sheet_view.rightToLeft = True
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 24
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                42, max(12, width + 2)
            )
            for cell in column[1:]:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = (
                        "dd/mm/yyyy hh:mm:ss.000"
                        if column[0].value == "scraped_at"
                        else "dd/mm/yyyy hh:mm"
                    )
                elif isinstance(cell.value, str):
                    # Store source strings literally, including any text beginning with '='.
                    cell.data_type = "s"
                elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                    cell.number_format = "0"
    workbook.properties.title = "Ligat Al data"
    workbook.properties.description = (
        "Eight clean datasets; canonical Hebrew names from teams.yaml; "
        "all dates use Asia/Jerusalem local time."
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.{uuid4().hex}.xlsx")
    metadata_temporary = temporary.with_suffix(".json.tmp")
    try:
        workbook.save(temporary)
        saved = load_workbook(temporary, read_only=True, data_only=False)
        try:
            if saved.sheetnames != list(DATASETS):
                raise ValueError("Unexpected Excel worksheet names")
            for name, frame in frames.items():
                sheet = saved[name]
                if sheet.max_row != len(frame) + 1 or sheet.max_column != len(frame.columns):
                    raise ValueError(f"Excel row/column count mismatch: {name}")
        finally:
            saved.close()
        if hashes is not None:
            if hashes != _fingerprints(repo, teams_path):
                raise ValueError("Parquet changed during export; rerun synchronization")
            # Fail on collision: a numbered version is never replaced.
            metadata = {
                "version": output.stem.removeprefix("ligat_al_data_"),
                "source_sha256": hashes,
                "excel_sha256": hashlib.sha256(temporary.read_bytes()).hexdigest(),
                "rows": {name: len(frame) for name, frame in frames.items()},
            }
            metadata_temporary.write_text(json.dumps(metadata, indent=2), encoding="utf-8")
            os.link(temporary, output)
            os.link(metadata_temporary, output.with_suffix(".json"))
        else:
            os.replace(temporary, output)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
        metadata_temporary.unlink(missing_ok=True)
    return output
