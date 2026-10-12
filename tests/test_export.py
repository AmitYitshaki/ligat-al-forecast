"""Workbook round trip, canonical names, versioning and immutable source checks."""

import hashlib
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from ligat.data.export import DATASETS, export_excel
from ligat.data.repository import Repository


@pytest.mark.parametrize("with_squad_values", [False, True])
def test_export_sheets_and_versioning(tmp_path: Path, with_squad_values: bool) -> None:
    teams = tmp_path / "teams.yaml"
    teams.write_text("teams:\n- team_id: 1\n  name_he: '=literal'\n", encoding="utf-8")
    repo = Repository(tmp_path / "clean")
    repo.root.mkdir()
    datasets = DATASETS + (("squad_values",) if with_squad_values else ())
    for name in datasets:
        (repo.root / f"{name}.parquet").write_bytes(name.encode())
    hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in repo.root.iterdir()}
    date = pd.Timestamp("2026-10-12 18:00", tz="Asia/Jerusalem")
    with patch.object(
        Repository,
        "read",
        side_effect=lambda _: pd.DataFrame({"team_id": [1], "date": [date], "missing": [None]}),
    ):
        output = export_excel(repo, teams_path=teams)
        assert output.name == "ligat_al_data_0.3.xlsx"
        assert export_excel(repo, teams_path=teams) == output
        book = load_workbook(output)
        assert book.sheetnames == list(datasets)
        for sheet in book:
            assert [cell.value for cell in sheet[1]] == [
                "team_id",
                "team_name_he",
                "date",
                "missing",
            ]
            assert sheet["B2"].value == "=literal" and sheet["B2"].data_type == "s"
            assert sheet["C2"].value == date.tz_localize(None).to_pydatetime()
            assert sheet.sheet_view.rightToLeft and sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref == "A1:D2"
        book.close()
        teams.write_text("teams:\n- team_id: 1\n  name_he: Updated\n", encoding="utf-8")
        assert export_excel(repo, teams_path=teams).name == "ligat_al_data_0.4.xlsx"
        with pytest.raises(FileExistsError):
            export_excel(repo, output=output, teams_path=teams)
    assert all(
        hashlib.sha256((repo.root / p).read_bytes()).hexdigest() == h for p, h in hashes.items()
    )
